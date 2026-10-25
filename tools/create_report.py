from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Assignment5_Strategy_Report_Amirkhan_Tulkibayev.docx"
DIAGRAM = ROOT / "docs" / "uml-class-diagram.png"
GITHUB_URL = "https://github.com/hann-qaz/Assignment_5_Strategy_Pattern"


def set_font(run, name="Arial", size=11, bold=False, italic=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color="666666")
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))


def add_paragraph(document, text="", bold_lead=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = paragraph.add_run(text)
        set_font(run)
    return paragraph


def add_code(document, code):
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F4F6")
    set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(code.strip("\n").splitlines()):
        run = paragraph.add_run(line)
        set_font(run, name="Courier New", size=8.5)
        if index < len(code.strip("\n").splitlines()) - 1:
            run.add_break()


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def add_before_after(document, before_code, after_code):
    before_label = add_paragraph(document, "Before", bold_lead="Before")
    before_label.paragraph_format.keep_with_next = True
    add_code(document, before_code)
    after_label = add_paragraph(document, "After", bold_lead="After")
    after_label.paragraph_format.keep_with_next = True
    add_code(document, after_code)


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)

for style_name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 12.5)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True

title_style_properties = styles["Title"]._element.get_or_add_pPr()
title_border = title_style_properties.find(qn("w:pBdr"))
if title_border is not None:
    title_style_properties.remove(title_border)

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
footer_run = footer_paragraph.add_run("Amirkhan Tulkibayev  |  SE-2522")
set_font(footer_run, size=9, color="666666")
add_page_number(footer.add_paragraph())

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(70)
title.paragraph_format.space_after = Pt(14)
title_run = title.add_run("Assignment 5 Strategy Pattern Report")
set_font(title_run, size=23, bold=True)

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(34)
subtitle_run = subtitle.add_run("Navigation Route Planning")
set_font(subtitle_run, size=15, bold=True)

for label, value in (
    ("Student", "Amirkhan Tulkibayev"),
    ("Group", "SE-2522"),
    ("Course", "Software Design Patterns"),
    ("Language", "Java 17"),
):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    label_run = paragraph.add_run(f"{label}  ")
    set_font(label_run, bold=True)
    value_run = paragraph.add_run(value)
    set_font(value_run)

document.add_paragraph()
summary = document.add_paragraph()
summary.paragraph_format.space_before = Pt(25)
summary.paragraph_format.line_spacing = 1.15
summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
summary_run = summary.add_run(
    "The project uses interchangeable route algorithms and demonstrates runtime strategy switching without transport-selection branches in the Context."
)
set_font(summary_run, size=11.5)

document.add_page_break()

add_heading(document, "1 Introduction")
add_paragraph(
    document,
    "This project applies the Strategy pattern to route planning. I implemented four interchangeable ways to plan the same trip: car, walking, bicycle, and public transport. Every algorithm receives the same RouteRequest and returns a RoutePlan, but each one calculates time with its own speed and extra delay. The client can replace the algorithm while the program is running.",
)

add_heading(document, "2 UML Class Diagram")
figure = document.add_paragraph()
figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure.paragraph_format.keep_with_next = True
figure.add_run().add_picture(str(DIAGRAM), width=Inches(5.35))
caption = document.add_paragraph()
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_after = Pt(8)
caption_run = caption.add_run("Figure 1  Strategy pattern structure")
set_font(caption_run, size=9, italic=True, color="555555")

for text in (
    "RouteStrategy is the common Strategy interface.",
    "CarRouteStrategy, WalkingRouteStrategy, BicycleRouteStrategy, and PublicTransportRouteStrategy are Concrete Strategies.",
    "Navigator is the Context. It stores a RouteStrategy reference and delegates route planning.",
    "NavigationDemo is the Client. It creates strategies and switches them at runtime.",
):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    set_font(paragraph.add_run(text))

document.add_page_break()

add_heading(document, "3 Clean Code")

add_heading(document, "3 1 Polymorphism Replaces Strategy Selection Conditionals", level=2)
add_before_after(
    document,
    '''switch (transportType) {
    case "car" -> calculateCarRoute(request);
    case "walk" -> calculateWalkingRoute(request);
    case "bicycle" -> calculateBicycleRoute(request);
}''',
    '''public RoutePlan buildRoute(RouteRequest request) {
    return routeStrategy.plan(request);
}''',
)
add_paragraph(
    document,
    "Navigator delegates through the common interface, so it has no conditional logic for choosing an algorithm. The if statements in RouteRequest only validate input; they do not select a strategy.",
)

add_heading(document, "3 2 Meaningful Names Reveal Intent", level=2)
add_before_after(
    document,
    "class TypeA { Result doIt(Data x) { ... } }",
    '''class WalkingRouteStrategy implements RouteStrategy {
    public RoutePlan plan(RouteRequest request) { ... }
}''',
)
add_paragraph(
    document,
    "WalkingRouteStrategy, RouteRequest, estimatedMinutes, and buildRoute explain the domain without extra comments.",
)

add_heading(document, "3 3 Each Class Has One Responsibility", level=2)
add_before_after(
    document,
    '''class Navigator {
    // car + walking + bicycle + bus formulas
}''',
    '''public class BicycleRouteStrategy implements RouteStrategy {
    public RoutePlan plan(RouteRequest request) {
        // only the bicycle calculation
    }
}''',
)
add_paragraph(
    document,
    "Each Concrete Strategy owns one route algorithm. Changing bicycle speed cannot affect the walking or car calculation.",
)

add_heading(document, "3 4 The Context Is Algorithm Agnostic", level=2)
add_before_after(
    document,
    "private final CarRouteStrategy strategy = new CarRouteStrategy();",
    '''private RouteStrategy routeStrategy;

public void setRouteStrategy(RouteStrategy routeStrategy) {
    this.routeStrategy = Objects.requireNonNull(routeStrategy);
}''',
)
add_paragraph(
    document,
    "Navigator depends only on the RouteStrategy interface. It does not know concrete strategy types or their formulas.",
)

add_heading(document, "3 5 The Design Is Open for Extension", level=2)
add_before_after(
    document,
    'case "bus" -> calculateBusRoute(request);',
    '''public class PublicTransportRouteStrategy implements RouteStrategy {
    public RoutePlan plan(RouteRequest request) { ... }
}''',
)
add_paragraph(
    document,
    "Public transport was added as a separate class. Navigator, RouteStrategy, and the existing strategies did not change. The Client only passes the new object to the Context.",
)

add_heading(document, "4 Strategy and State")
add_paragraph(
    document,
    "The course comparison asks a useful question. Strategy answers which algorithm should perform a task, while State answers which lifecycle stage an object is in. In Strategy, the caller or configuration normally chooses the behavior. In State, the Context or a state object usually controls transitions and changes how events are handled.",
)
add_paragraph(
    document,
    "This navigation problem needs Strategy because car, walking, bicycle, and public transport are alternative algorithms for one trip. They are not ordered stages, there are no transition rules, and the caller may choose any route method at any time. Runtime switching alone does not make the design State.",
)

add_heading(document, "5 Conclusion")
add_paragraph(
    document,
    "The Strategy pattern adds several small classes, but it makes the design flexible and easy to test. Navigator stays simple, each algorithm can change independently, and another route type can be added without modifying existing classes. The extra classes are a reasonable trade-off for clear responsibilities and safe extension.",
)

add_heading(document, "6 GitHub Repository")
link_paragraph = document.add_paragraph()
add_hyperlink(link_paragraph, GITHUB_URL, GITHUB_URL)

add_heading(document, "References")
references = (
    "Ospanov, Almas. Lecture 7 The Strategy Pattern. Astana IT University, 2026.",
    "Bekenov, Yerassyl. Strategy vs State Untangling Two Look-Alike Design Patterns. Astana IT University, 20 October 2026.",
)
for index, item in enumerate(references, start=1):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_font(paragraph.add_run(f"{index}. {item}"))

document.core_properties.title = "Assignment 5 Strategy Pattern Report"
document.core_properties.author = "Amirkhan Tulkibayev"
document.core_properties.subject = "Strategy Pattern navigation route planning"
document.save(OUTPUT)
print(OUTPUT)

